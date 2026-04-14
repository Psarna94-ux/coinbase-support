import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Account Flags"

# Styles
header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="0052FF", end_color="0052FF", fill_type="solid")
wrap = Alignment(wrap_text=True, vertical="top")
thin_border = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)
category_font = Font(name="Calibri", bold=True, size=10, color="0A0B0D")
category_fill = PatternFill(start_color="EEF0F3", end_color="EEF0F3", fill_type="solid")

# Column widths
ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 35
ws.column_dimensions["D"].width = 55
ws.column_dimensions["E"].width = 65
ws.column_dimensions["F"].width = 18

# Headers
headers = ["#", "Category", "Flag Name", "Description / Reason", "SoP to Resolve", "Severity"]
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"

flags = [
    # KYC / Identity Verification
    ("KYC / Identity", "KYC_INCOMPLETE", "User has not completed mandatory Know Your Customer identity verification.", "Send automated reminder via email/push. If unresolved after 30 days, restrict trading. Escalate to L2 if user disputes. Provide self-serve link to upload documents.", "Medium"),
    ("KYC / Identity", "KYC_DOC_EXPIRED", "Government-issued ID submitted during verification has expired.", "Notify user to upload a valid, non-expired ID. Temporarily restrict withdrawals over $1,000 until resolved. Auto-clear flag upon successful re-verification.", "Medium"),
    ("KYC / Identity", "KYC_DOC_MISMATCH", "Name or date of birth on submitted ID does not match account registration details.", "Restrict account trading. Request user to submit corrected documentation or legal name change proof. Escalate to compliance analyst for manual review.", "High"),
    ("KYC / Identity", "KYC_SELFIE_FAIL", "Liveness/selfie verification failed — photo does not match the submitted ID.", "Allow up to 3 retry attempts within 24 hours. If all fail, escalate to manual review by identity verification team. Suspend deposits until cleared.", "High"),
    ("KYC / Identity", "KYC_PEP_MATCH", "User matched against Politically Exposed Persons (PEP) database.", "Route to Enhanced Due Diligence (EDD) queue. Compliance analyst must review source of funds, nature of business, and approve/deny within 5 business days. Document decision in case file.", "Critical"),
    ("KYC / Identity", "KYC_SANCTIONS_HIT", "User name or address matched against OFAC/SDN or other sanctions lists.", "Immediately freeze all account activity. Escalate to BSA/AML Officer. File SAR if confirmed. Do not notify user of the sanctions match. Legal review required before any action.", "Critical"),
    ("KYC / Identity", "KYC_ADVERSE_MEDIA", "Negative media screening flagged the user in connection with financial crime, fraud, or terrorism.", "Route to EDD team. Analyst reviews media sources for credibility and relevance. If confirmed, restrict account and file SAR. If false positive, document and clear.", "High"),
    ("KYC / Identity", "KYC_UNDERAGE", "User's date of birth indicates they are under 18 years of age.", "Immediately restrict account. Notify user that Coinbase requires users to be 18+. Initiate account closure and refund any balances to original payment method.", "High"),

    # Transaction Monitoring
    ("Transaction Monitoring", "TXN_HIGH_VELOCITY", "Unusually high number of transactions in a short time window (e.g., >50 txns in 1 hour).", "Auto-restrict outgoing transfers. Alert L1 analyst for review within 2 hours. If legitimate (e.g., API trader), whitelist after verification. Otherwise escalate for SAR consideration.", "High"),
    ("Transaction Monitoring", "TXN_LARGE_DEPOSIT", "Single deposit exceeding $10,000 or cumulative deposits exceeding $25,000 in 24 hours.", "Trigger CTR filing if cash-equivalent. Hold funds for 24-hour review. Analyst verifies source of funds documentation. Release upon approval or escalate.", "Medium"),
    ("Transaction Monitoring", "TXN_STRUCTURING", "Multiple deposits just below reporting thresholds suggesting deliberate structuring to evade CTR filing.", "Immediately escalate to BSA/AML team. Restrict further deposits. File SAR within 30 days. Do not inform user of the structuring suspicion.", "Critical"),
    ("Transaction Monitoring", "TXN_RAPID_WITHDRAWAL", "User deposited funds and attempted full withdrawal within 24 hours.", "Hold withdrawal for manual review. Verify deposit has fully cleared. Check for chargeback risk. Analyst approves or denies within 4 hours.", "High"),
    ("Transaction Monitoring", "TXN_MIXER_DETECTED", "Funds traced to or from a known cryptocurrency mixing/tumbling service.", "Restrict outgoing transfers. Escalate to blockchain analytics team for full trace. If confirmed mixer interaction, file SAR and consider account termination.", "Critical"),
    ("Transaction Monitoring", "TXN_DARKNET_EXPOSURE", "Blockchain analysis indicates funds originated from or were sent to a darknet marketplace.", "Immediately freeze account. Escalate to BSA/AML Officer and legal team. File SAR. Coordinate with law enforcement if subpoena received.", "Critical"),
    ("Transaction Monitoring", "TXN_CROSS_BORDER_RISK", "High-value transfers to/from jurisdictions on FATF grey/black list.", "Route to compliance review. Analyst verifies purpose of transaction and beneficiary details. May require additional documentation from user. Apply enhanced monitoring.", "High"),
    ("Transaction Monitoring", "TXN_ROUND_AMOUNTS", "Repeated transactions in suspiciously round amounts (e.g., $5,000, $10,000) with no clear business rationale.", "Flag for analyst review. Compare against user's stated income and transaction history. If inconsistent, escalate for SAR filing.", "Medium"),

    # Account Security
    ("Account Security", "SEC_ACCOUNT_TAKEOVER", "Indicators of unauthorized account access — new device, IP change, immediate withdrawal attempt.", "Immediately lock account. Send security alert to registered email/phone. Require full re-authentication including ID verification. Review last 72 hours of activity.", "Critical"),
    ("Account Security", "SEC_SIM_SWAP_RISK", "Phone number recently ported or SIM swap detected for account with SMS-based 2FA.", "Lock account and disable SMS 2FA. Notify user via email. Require in-app identity re-verification before restoring access. Recommend hardware key or authenticator app.", "Critical"),
    ("Account Security", "SEC_CREDENTIAL_LEAK", "User's email/password found in a known data breach database.", "Force password reset on next login. Send breach notification email. Require 2FA setup if not already enabled. Monitor for unauthorized access for 30 days.", "High"),
    ("Account Security", "SEC_MULTIPLE_DEVICES", "Account accessed from more than 5 unique devices in 7 days.", "Trigger step-up authentication. Send notification listing recent devices. Allow user to revoke unrecognized sessions. Monitor for 14 days.", "Medium"),
    ("Account Security", "SEC_VPN_TOR_ACCESS", "Account accessed via VPN or Tor exit node from a restricted jurisdiction.", "Temporarily restrict trading. Prompt user to verify location. If user is in a restricted jurisdiction, initiate account wind-down. Log for compliance records.", "High"),
    ("Account Security", "SEC_PHISHING_VICTIM", "User reported falling victim to phishing — credentials may be compromised.", "Immediately lock account. Reset all credentials. Review transaction history for unauthorized activity. Assist user with recovery process. File incident report.", "Critical"),

    # Payment & Chargeback
    ("Payment / Chargeback", "PAY_CHARGEBACK_FILED", "User's bank initiated a chargeback on a deposit that was already traded or withdrawn.", "Restrict account trading and withdrawals. Calculate outstanding negative balance. Send repayment notice to user. If unresolved in 30 days, send to collections.", "High"),
    ("Payment / Chargeback", "PAY_NEGATIVE_BALANCE", "Account has a negative fiat balance due to reversed deposit, failed payment, or fee adjustment.", "Restrict all activity until balance is repaid. Send automated repayment reminders at 7, 14, and 30 days. Offer partial payment plans for balances over $500.", "High"),
    ("Payment / Chargeback", "PAY_FRAUD_DEPOSIT", "Deposit made with a stolen credit card or fraudulent bank account.", "Immediately freeze deposited funds. Reverse any trades made with fraudulent funds. File SAR. Coordinate with payment processor for fraud investigation.", "Critical"),
    ("Payment / Chargeback", "PAY_MULTIPLE_CARDS", "More than 5 different payment cards linked to the account within 30 days.", "Restrict new card additions. Analyst reviews card ownership verification. If cards belong to different individuals, escalate for fraud investigation.", "Medium"),
    ("Payment / Chargeback", "PAY_BANK_RETURN", "ACH/bank transfer returned due to insufficient funds, closed account, or unauthorized transaction.", "Apply returned deposit fee. Restrict ACH deposits for 14 days. If repeated (3+ times), permanently disable bank transfer deposits.", "Medium"),
    ("Payment / Chargeback", "PAY_FAILED_SETTLEMENT", "Fiat settlement to user's bank failed due to incorrect account details or bank rejection.", "Notify user to update bank details. Hold funds in Coinbase balance. Retry settlement after user confirms correct details. Escalate after 3 failed attempts.", "Low"),

    # Fraud & Abuse
    ("Fraud / Abuse", "FRD_MULTI_ACCOUNT", "User operating multiple accounts in violation of Coinbase's one-account policy.", "Identify all linked accounts via email, phone, device fingerprint, IP. Restrict secondary accounts. Merge legitimate activity into primary. Warn user of TOS violation.", "High"),
    ("Fraud / Abuse", "FRD_PROMO_ABUSE", "User exploiting referral program, sign-up bonuses, or promotional campaigns through multiple accounts or fake referrals.", "Claw back promotional rewards. Restrict referral program access. If systematic, escalate to fraud team and consider account termination.", "Medium"),
    ("Fraud / Abuse", "FRD_IDENTITY_THEFT", "Evidence suggests the account was opened using stolen identity documents.", "Immediately freeze account. Initiate identity theft investigation. Request additional verification from account holder. File SAR and coordinate with ID theft victim if identified.", "Critical"),
    ("Fraud / Abuse", "FRD_SCAM_VICTIM", "User appears to be victim of an investment scam, romance scam, or social engineering attack.", "Temporarily restrict large outgoing transfers. Contact user with scam awareness information. Offer to reverse recent transactions if possible. Document interaction.", "High"),
    ("Fraud / Abuse", "FRD_MULE_ACCOUNT", "Account behavior consistent with a money mule — receiving and quickly forwarding funds for third parties.", "Restrict all outgoing transfers. Escalate to financial crimes team. File SAR. Consider account termination. Preserve all records for potential law enforcement request.", "Critical"),
    ("Fraud / Abuse", "FRD_WASH_TRADING", "User engaging in self-trading or coordinated trading to manipulate market prices or volume.", "Restrict trading on affected pairs. Escalate to market surveillance team. If confirmed, claw back ill-gotten gains and issue TOS violation warning.", "High"),
    ("Fraud / Abuse", "FRD_INSIDER_TRADING", "Trading patterns suggest user acted on material non-public information about token listings or delistings.", "Immediately restrict account. Escalate to legal and compliance leadership. Preserve all trade records. Cooperate with SEC/DOJ if investigation initiated.", "Critical"),

    # Regulatory / Compliance
    ("Regulatory", "REG_CTR_REQUIRED", "Cash or cash-equivalent transaction of $10,000+ requires Currency Transaction Report filing.", "Auto-generate CTR. Compliance analyst reviews and files with FinCEN within 15 days. No user notification required. Flag auto-clears after filing.", "Medium"),
    ("Regulatory", "REG_SAR_FILED", "Suspicious Activity Report has been filed on this account.", "Maintain enhanced monitoring for 90 days. Do not disclose SAR filing to user (tipping off is a federal crime). Review account for continued suspicious activity.", "Critical"),
    ("Regulatory", "REG_TRAVEL_RULE", "Crypto transfer of $3,000+ requires Travel Rule compliance — originator/beneficiary information must be collected.", "Hold outgoing transfer until required originator and beneficiary information is collected and verified. Auto-release upon successful VASP-to-VASP data exchange.", "Medium"),
    ("Regulatory", "REG_TAX_HOLD", "IRS or state tax authority has issued a levy or hold on user's account.", "Restrict withdrawals per levy amount. Notify user of the legal hold. Coordinate with legal team to comply with levy requirements. Release only upon court order or IRS release.", "High"),
    ("Regulatory", "REG_SUBPOENA_HOLD", "Law enforcement subpoena or preservation request received for this account.", "Preserve all account data and transaction records. Do not notify user unless legally permitted. Restrict account modifications. Legal team manages response timeline.", "Critical"),
    ("Regulatory", "REG_JURISDICTION_BAN", "User's verified location is in a jurisdiction where Coinbase services are prohibited.", "Restrict all trading and deposits. Notify user they have 30 days to withdraw existing balances. Initiate orderly account wind-down. Block future sign-ins from that jurisdiction.", "High"),

    # Risk Scoring
    ("Risk Scoring", "RISK_HIGH_SCORE", "User's aggregate risk score exceeds threshold based on combined behavioral, transactional, and identity signals.", "Apply enhanced monitoring — lower withdrawal limits, require additional verification for large trades. Analyst reviews monthly. Adjust score as new data emerges.", "High"),
    ("Risk Scoring", "RISK_ESCALATED_REVIEW", "Account flagged for periodic enhanced due diligence review (e.g., high-net-worth, business account).", "Compliance analyst conducts full account review — transaction patterns, source of funds, account purpose. Update risk assessment. Schedule next review in 6–12 months.", "Medium"),
    ("Risk Scoring", "RISK_DORMANT_REACTIVATION", "Previously dormant account (no activity for 12+ months) suddenly shows high-value transactions.", "Trigger step-up authentication. Analyst reviews recent activity against historical patterns. If consistent, clear flag. If anomalous, restrict and investigate.", "Medium"),
    ("Risk Scoring", "RISK_BEHAVIORAL_ANOMALY", "Machine learning model detected transaction patterns significantly deviating from user's established behavior.", "Auto-restrict outgoing transfers over $1,000. Route to L1 analyst for review within 4 hours. If legitimate behavior change, update user profile and clear.", "Medium"),

    # Account Administrative
    ("Account Admin", "ADM_ACCOUNT_LOCKED", "Account manually locked by support agent due to user request or security concern.", "Verify user identity through full re-authentication. Review reason for lock. Remove lock only after identity confirmed and security concern addressed. Document in case notes.", "Low"),
    ("Account Admin", "ADM_DECEASED_USER", "Notification received that the account holder is deceased.", "Restrict all account activity. Request death certificate and estate documentation. Coordinate with estate executor for asset transfer. Legal team reviews and approves disbursement.", "High"),
    ("Account Admin", "ADM_LEGAL_NAME_CHANGE", "User requested account name change due to marriage, divorce, or court order.", "Request legal documentation (marriage certificate, court order). Compliance verifies documents. Update account name and re-run sanctions screening. Notify user of completion.", "Low"),
    ("Account Admin", "ADM_BUSINESS_CONVERSION", "Individual account flagged for business-level transaction volumes requiring entity account upgrade.", "Notify user of required upgrade to business account. Collect business documentation (EIN, articles of incorporation, beneficial ownership). Restrict individual account if not converted within 60 days.", "Medium"),
    ("Account Admin", "ADM_RESTRICTED_ASSET", "User holds or attempted to trade an asset that has been delisted or restricted in their jurisdiction.", "Notify user of the restriction with 14-day sell-only window. Disable buy orders for the asset. After grace period, auto-convert to USDC if user hasn't acted. Document regulatory basis for restriction.", "Medium"),
]

row = 2
current_category = ""
for cat, flag, desc, sop, severity in flags:
    ws.cell(row=row, column=1, value=row - 1).alignment = wrap
    ws.cell(row=row, column=1, value=row - 1).border = thin_border

    cat_cell = ws.cell(row=row, column=2, value=cat)
    cat_cell.alignment = wrap
    cat_cell.border = thin_border
    if cat != current_category:
        cat_cell.font = category_font
        current_category = cat

    flag_cell = ws.cell(row=row, column=3, value=flag)
    flag_cell.alignment = wrap
    flag_cell.border = thin_border
    flag_cell.font = Font(name="Consolas", size=10)

    desc_cell = ws.cell(row=row, column=4, value=desc)
    desc_cell.alignment = wrap
    desc_cell.border = thin_border

    sop_cell = ws.cell(row=row, column=5, value=sop)
    sop_cell.alignment = wrap
    sop_cell.border = thin_border

    sev_cell = ws.cell(row=row, column=6, value=severity)
    sev_cell.alignment = Alignment(horizontal="center", vertical="top")
    sev_cell.border = thin_border

    # Color-code severity
    if severity == "Critical":
        sev_cell.font = Font(bold=True, color="FFFFFF", size=10)
        sev_cell.fill = PatternFill(start_color="CF202F", end_color="CF202F", fill_type="solid")
    elif severity == "High":
        sev_cell.font = Font(bold=True, color="FFFFFF", size=10)
        sev_cell.fill = PatternFill(start_color="E67E22", end_color="E67E22", fill_type="solid")
    elif severity == "Medium":
        sev_cell.font = Font(bold=True, color="0A0B0D", size=10)
        sev_cell.fill = PatternFill(start_color="F1C40F", end_color="F1C40F", fill_type="solid")
    elif severity == "Low":
        sev_cell.font = Font(bold=True, color="FFFFFF", size=10)
        sev_cell.fill = PatternFill(start_color="3CC28A", end_color="3CC28A", fill_type="solid")

    ws.row_dimensions[row].height = 80
    row += 1

wb.save("/Users/pranusarna/coinbase-support/Coinbase_Account_Flags.xlsx")
print(f"Done — {row - 2} flags written")
